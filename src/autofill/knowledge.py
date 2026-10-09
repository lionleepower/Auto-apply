"""知识库：把用户随手写的文字 / 上传的简历整理成 Profile，并指出还缺什么。

文字先在本地抽取（PDF / Word / 文本），只有用户点「AI 整理」时才发给所选模型；
身份证号这类证件号发送前先打码，也永远不会被自动填写。
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

import yaml

from .llm import ChatClient
from .profile import (Basic, Education, Experience, Paper, Profile, Project, flatten, parse_doc,
                      dump_doc)

# ---------- 文本提取 ----------

MAX_CHARS = 30000


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        text = "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages)
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # 很多简历模板用表格排版
            for row in table.rows:
                seen: list[str] = []
                for cell in row.cells:
                    if cell.text not in seen:  # 合并单元格会重复出现
                        seen.append(cell.text)
                parts.append(" | ".join(seen))
        text = "\n".join(parts)
    elif ext in (".txt", ".md", ".text", ""):
        for enc in ("utf-8-sig", "gbk"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"暂不支持 {ext} 格式，请上传 PDF / Word(.docx) / 文本，或把内容粘贴进知识库")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("没读到文字：可能是扫描件 / 图片版 PDF。请换文字版，或把内容粘贴进知识库")
    return text[:MAX_CHARS]


PORTFOLIO_EXTS = (".zip", ".pdf", ".rar", ".7z")
TEXT_NAMES = re.compile(r"(^|/)(readme[^/]*|[^/]+\.(md|txt|rst))$", re.I)


def extract_zip_text(data: bytes, max_chars: int = 12000) -> str:
    """作品集 zip 里的 README / 文档文字（不解压到磁盘，只读文本类小文件，防压缩炸弹）。读不出来返回空串。"""
    import zipfile

    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise ValueError("这不是有效的 zip 文件") from e
    out, total = [], 0
    infos = sorted((i for i in zf.infolist() if not i.is_dir() and TEXT_NAMES.search(i.filename)
                    and ".." not in i.filename and not i.filename.startswith("/")),
                   key=lambda i: (i.filename.count("/"), i.filename))  # 越靠近根目录越先读
    for info in infos[:12]:
        if info.file_size > 200_000:
            continue
        try:
            raw = zf.read(info)
        except Exception:  # 加密或损坏的条目：跳过
            continue
        for enc in ("utf-8-sig", "gbk"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            continue
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if not text:
            continue
        chunk = f"—— 作品集文件 {info.filename} ——\n{text}"
        out.append(chunk[: max_chars - total])
        total += len(out[-1])
        if total >= max_chars:
            break
    return "\n\n".join(out)


ID_RE = re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")


def redact(text: str) -> str:
    return ID_RE.sub("[证件号已隐藏]", text)


# ---------- 缺失项 ----------

# (键, 名称, 级别)：required 几乎每张表都要；recommended 不少表会问
CHECKS_CN: list[tuple[str, str, str]] = [
    ("basic.name", "姓名", "required"),
    ("basic.phone", "手机号", "required"),
    ("basic.email", "邮箱", "required"),
    ("basic.gender", "性别", "required"),
    ("basic.birth_date", "出生日期", "required"),
    ("education.school", "学校", "required"),
    ("education.major", "专业", "required"),
    ("education.degree", "学历", "required"),
    ("education.start", "入学时间", "required"),
    ("education.end", "毕业时间", "required"),
    ("basic.hometown", "籍贯 / 生源地", "recommended"),
    ("basic.current_city", "现居城市", "recommended"),
    ("basic.expected_city", "期望工作城市", "recommended"),
    ("basic.political_status", "政治面貌", "recommended"),
    ("basic.available", "到岗时间", "recommended"),
    ("education.gpa", "GPA / 成绩", "recommended"),
    ("skills", "专业技能", "recommended"),
    ("self_intro", "自我评价", "recommended"),
]
CHECKS_OVERSEAS: list[tuple[str, str, str]] = [
    ("basic.name_en", "英文姓名", "required"),
    ("basic.email", "邮箱", "required"),
    ("basic.phone", "手机号", "required"),
    ("basic.country", "现居国家 / 地区", "required"),
    ("basic.authorized_to_work", "是否有当地合法工作资格", "required"),
    ("basic.need_sponsorship", "是否需要签证担保", "required"),
    ("education.school", "学校", "required"),
    ("education.major", "专业", "required"),
    ("education.degree", "学历", "required"),
    ("education.start", "入学时间", "required"),
    ("education.end", "毕业时间", "required"),
    ("basic.linkedin", "LinkedIn", "recommended"),
    ("basic.address", "街道地址", "recommended"),
    ("basic.current_city", "现居城市", "recommended"),
    ("basic.available", "到岗时间", "recommended"),
    ("basic.work_authorization", "工作许可状态", "recommended"),
    ("education.gpa", "GPA / 成绩", "recommended"),
    ("skills", "专业技能", "recommended"),
    ("self_intro", "自我介绍 / Summary", "recommended"),
]
SECTION_OF = {"basic": "basic", "education": "education", "skills": "extra", "self_intro": "extra"}


def checks_for(profile: Profile) -> list[tuple[str, str, str]]:
    return CHECKS_OVERSEAS if profile.locale == "overseas" else CHECKS_CN


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _value(profile: Profile, key: str):
    v = profile.get(key, 0)
    if key == "basic.name_en" and _blank(v) and profile.basic.first_name and profile.basic.last_name:
        return f"{profile.basic.first_name} {profile.basic.last_name}"
    return v


def missing(profile: Profile) -> list[dict]:
    out = []
    for key, label, level in checks_for(profile):
        head = key.split(".")[0]
        if head == "education":
            # 第一段教育经历（最高学历）必须有；缺整段时只报一次
            if not profile.education:
                if key == "education.school":
                    out.append({"key": "education", "label": "教育经历", "level": level, "section": "education"})
                continue
            if _blank(profile.get(key, 0)):
                out.append({"key": key, "label": f"教育经历 · {label}", "level": level, "section": "education"})
        elif _blank(_value(profile, key)):
            out.append({"key": key, "label": label, "level": level, "section": SECTION_OF.get(head, "basic")})
    if not (profile.experience or profile.projects):
        out.append({"key": "experience", "label": "实习 / 项目经历（至少一段）", "level": "recommended",
                    "section": "experience"})
    return out


def completeness(profile: Profile) -> int:
    checks = checks_for(profile)
    filled = sum(not _blank(_value(profile, key)) for key, _, _ in checks)
    filled += bool(profile.experience or profile.projects)
    return round(100 * filled / (len(checks) + 1))


# ---------- 抽取 ----------

EXTRACT_SYSTEM = """你是简历信息抽取器。用户给你一段简历或个人信息文字，把它整理成固定结构的 JSON。
规则：
1. 只抽取文字里明确出现的信息，没有的字段一律用 null，绝不编造、不推测。
2. 日期：出生日期 YYYY-MM-DD；经历起止时间 YYYY-MM；至今 / 在读的结束时间写 null。
3. degree 只能是 本科 / 硕士 / 博士 / 专科 之一；gender 只能是 男 / 女。
4. 教育经历、实习工作、项目各按时间倒序（最近的在前）。experience.kind 为 实习 / 工作 / 科研 之一。
5. description 保留原文要点，可合并成一段，不要扩写。
6. 证件号不要抽取。
7. 保持简历原文的语言，不要翻译（英文简历抽出来的就是英文）。
只输出 JSON，结构如下（缺的字段写 null，列表没有就写 []）：
{"basic": {"name","name_en","gender","birth_date","phone","email","political_status","ethnicity","hometown",
"current_city","expected_city","job_intention","available","english","github","homepage",
"first_name","last_name","linkedin","address","country","state","postal_code","nationality"},
"education": [{"school","college","major","degree","start","end","gpa","rank","city","advisor","research_area"}],
"experience": [{"kind","company","title","department","city","start","end","description"}],
"projects": [{"name","role","start","end","tech","url","description"}],
"papers": [{"title","venue","year","author_order"}],
"skills": "专业技能，一段文字", "awards": "获奖情况，一段文字", "self_intro": "自我评价"}"""


def _clean(model_cls, data) -> dict:
    """只保留模型认识的字段，值统一转成字符串（bool 字段保持 bool），丢掉空值。"""
    if not isinstance(data, dict):
        return {}
    out = {}
    for name, field in model_cls.model_fields.items():
        v = data.get(name)
        if _blank(v):
            continue
        if "bool" in str(field.annotation):
            out[name] = v if isinstance(v, bool) else str(v).strip() in ("是", "true", "True", "1")
        else:
            out[name] = str(v).strip()
    return out


def _text_field(v) -> str | None:
    if _blank(v):
        return None
    return "、".join(map(str, v)) if isinstance(v, list) else str(v).strip()


def parse_extraction(text: str) -> Profile:
    """解析模型输出。坏 JSON 抛 ValueError；个别字段不合法就丢掉，不让整份作废。"""
    raw = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.M).strip()
    start, end = raw.find("{"), raw.rfind("}")
    if start < 0 or end < start:
        raise ValueError("模型没有返回 JSON")
    try:
        data = json.loads(raw[start:end + 1])
    except json.JSONDecodeError as e:
        raise ValueError(f"模型返回的 JSON 无法解析：{e}") from e
    if not isinstance(data, dict):
        raise ValueError("模型返回的 JSON 不是对象")
    basic = _clean(Basic, data.get("basic"))
    basic.pop("id_number", None)

    def items(cls, key):
        return [cls(**c) for x in data.get(key) or [] if (c := _clean(cls, x))]

    return Profile(
        basic=Basic(**basic),
        education=items(Education, "education"),
        experience=items(Experience, "experience"),
        projects=items(Project, "projects"),
        papers=items(Paper, "papers"),
        skills=_text_field(data.get("skills")),
        awards=_text_field(data.get("awards")),
        self_intro=_text_field(data.get("self_intro")),
    )


def extract_with_llm(client: ChatClient, text: str) -> Profile:
    return parse_extraction(client.complete(EXTRACT_SYSTEM, redact(text)))


def extract_quick(text: str) -> Profile:
    """没配 AI 时的保底：只认手机号、邮箱、GitHub 这类格式固定的信息。"""
    basic: dict = {}
    if m := re.search(r"(?<!\d)1[3-9]\d{9}(?!\d)", text):
        basic["phone"] = m.group()
    if m := re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text):
        basic["email"] = m.group()
    if m := re.search(r"(?:https?://)?github\.com/[\w.-]+", text, re.I):
        basic["github"] = m.group()
    if m := re.search(r"(?:姓名|名字)\s*[:：]\s*([一-龥·]{2,5})", text):
        basic["name"] = m.group(1)
    if m := re.search(r"性别\s*[:：]\s*([男女])", text):
        basic["gender"] = m.group(1)
    if m := re.search(r"出生(?:日期|年月)?\s*[:：]\s*(\d{4})[-/.年](\d{1,2})(?:[-/.月](\d{1,2}))?", text):
        basic["birth_date"] = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3) or 1):02d}"
    return Profile(basic=Basic(**basic))


# ---------- 合并 ----------

SECTION_NAMES = {"education": "教育经历", "experience": "经历", "projects": "项目", "papers": "论文"}
KEY_ATTRS = {"education": ("school",), "experience": ("company", "title"), "projects": ("name",),
             "papers": ("title",)}


def _key(section: str, item) -> str:
    return "|".join((getattr(item, a) or "").strip().lower() for a in KEY_ATTRS[section])


def _fill(target, source, overwrite: bool) -> list[str]:
    changed = []
    for name in type(source).model_fields:
        nv = getattr(source, name)
        if _blank(nv):
            continue
        ov = getattr(target, name)
        if _blank(ov) or (overwrite and ov != nv):
            setattr(target, name, nv)
            changed.append(name)
    return changed


def merge(old: Profile, new: Profile, overwrite: bool = False) -> tuple[Profile, list[str]]:
    """把 new 并进 old。默认只补空：用户手改过的值不会被模型覆盖。返回（合并后的资料, 变更说明）。"""
    merged = old.model_copy(deep=True)
    notes = [f"补充 {n}" for n in _fill(merged.basic, new.basic, overwrite)]
    for section in KEY_ATTRS:
        have = {_key(section, x): x for x in getattr(merged, section)}
        for item in getattr(new, section):
            k = _key(section, item)
            if k in have:  # 已有同一段：只补它的空字段
                _fill(have[k], item, overwrite)
            else:
                getattr(merged, section).append(item)
                notes.append(f"新增{SECTION_NAMES[section]} · {k.split('|')[0]}")
    for k in ("skills", "awards", "self_intro"):
        nv = getattr(new, k)
        if not _blank(nv) and (_blank(getattr(merged, k)) or overwrite):
            setattr(merged, k, nv)
            notes.append(f"补充 {k}")
    return merged, notes


# ---------- 存取 ----------

def load_doc(path: Path) -> Profile:
    return parse_doc(yaml.safe_load(path.read_text(encoding="utf-8")) if path.exists() else {})


def load(path: Path, resume: str | None = None) -> Profile:
    """平铺视图（基本信息 + 某份简历），不补推导字段。"""
    return flatten(load_doc(path), resume, derive=False)


def save_doc(path: Path, doc: Profile) -> None:
    """写回 YAML。覆盖前备份成 .bak（原文件里的注释写回时会丢）。"""
    if path.exists():
        path.with_suffix(path.suffix + ".bak").write_bytes(path.read_bytes())
    path.write_text(yaml.safe_dump(dump_doc(doc), allow_unicode=True, sort_keys=False, width=100), encoding="utf-8")
